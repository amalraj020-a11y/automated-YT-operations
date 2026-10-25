"""Create a portable Excel report from one Amala video run."""
from __future__ import annotations

import argparse
import json
import os
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def _find(root: Path, name: str) -> Path | None:
    matches = sorted(root.rglob(name), key=lambda p: p.stat().st_mtime, reverse=True)
    return matches[0] if matches else None


def build_report(output_root: Path, workbook_path: Path) -> Path:
    result_path = _find(output_root, "result.json")
    content_path = _find(output_root, "content.json")
    if not result_path or not content_path:
        raise FileNotFoundError("Run output must contain result.json and content.json")
    result = json.loads(result_path.read_text(encoding="utf-8"))
    content = json.loads(content_path.read_text(encoding="utf-8"))
    packaging = content.get("packaging", {})
    hook = packaging.get("hook", {})
    title = packaging.get("title", {})
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    summary = book.active
    summary.title = "Run Summary"
    rows = [
        ("Generated (UTC)", datetime.now(timezone.utc).isoformat()),
        ("GitHub run", os.environ.get("GITHUB_RUN_ID", "")),
        ("Slot", result.get("slot", "")),
        ("Model", result.get("model", "")),
        ("Renderer", result.get("renderer", "")),
        ("Published", result.get("published", False)),
        ("YouTube URL", result.get("url", "")),
        ("Video file", result.get("video", "")),
        ("Title", content.get("title", "")),
        ("Title score", title.get("score", "")),
        ("Selected hook", hook.get("hook", "")),
        ("Hook score", hook.get("score", "")),
        ("Hook formula", hook.get("formula", "")),
    ]
    for row in rows:
        summary.append(row)
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 110

    content_sheet = book.create_sheet("Content")
    content_sheet.append(("Field", "Value"))
    for key in ("topic", "title", "description", "hashtags", "narration"):
        value = content.get(key, "")
        if isinstance(value, list):
            value = ", ".join(map(str, value))
        content_sheet.append((key, value))
    content_sheet.column_dimensions["A"].width = 20
    content_sheet.column_dimensions["B"].width = 120

    packaging_sheet = book.create_sheet("Packaging")
    packaging_sheet.append(("Type", "Text", "Score", "Formula", "Issues"))
    packaging_sheet.append(("Hook", hook.get("hook", ""), hook.get("score", ""), hook.get("formula", ""), ""))
    packaging_sheet.append(("Title", title.get("title", content.get("title", "")), title.get("score", ""), "", "; ".join(str(x) for x in title.get("issues", []))))
    packaging_sheet.column_dimensions["A"].width = 12
    packaging_sheet.column_dimensions["B"].width = 110
    packaging_sheet.column_dimensions["C"].width = 12
    for sheet in book.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        sheet.freeze_panes = "A2"
    book.save(workbook_path)
    return workbook_path


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--workbook", type=Path, required=True)
    args = parser.parse_args()
    print(build_report(args.output, args.workbook))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
